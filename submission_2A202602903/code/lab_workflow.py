"""Notebook helpers: diagnostics, controlled configs, official scoring and report."""
import json
import platform
import subprocess
import sys
from pathlib import Path
import numpy as np
import torch
import matplotlib.pyplot as plt
from model import MLP, count_params, EXPECTED_PARAMS
from train import DEFAULT_CFG, set_seed, compute_loss, evaluate, run_experiment, final_eval
from plots import plot_run, plot_compare
from results_table import save_result, to_row, write_xlsx

PREDICTIONS = {
 'baseline':'SGD+momentum .05 phải vượt mốc đa số; seed thay đổi thứ tự lô nên điểm có nhiễu.',
 'optimizer':'Adam cập nhật thích nghi có thể hội tụ nhanh hơn; mỗi optimizer cần thử nhiều lr.',
 'loss':'MSE logit/one-hot dự kiến học chậm hơn CE vì không trực tiếp tối ưu log-likelihood.',
 'hparam':'lr cao .5 có thể làm loss/gradient dao động hơn .05; chưa chắc phân kỳ.',
 'dropout':'q=.2 có thể giảm gap train–val nhưng làm chậm học nếu baseline chưa quá khớp.',
 'clipping':'Cắt ở phân vị 90% gradient baseline; có thể giảm gai ở lr cao, không chắc tăng F1.',
 'amp':'BF16 CPU có thể nhanh hơn nếu CPU hỗ trợ; F1 có thể lệch do làm tròn.',
 'init':'Xavier có phương sai nhỏ hơn He; zeros giữ đối xứng và chỉ học bias đầu ra.'}
MECHANISMS = {
 'baseline':'Seed tác động khởi tạo và thứ tự lô. σ từ ba lần chỉ là ước lượng thô.',
 'optimizer':'Momentum tích luỹ hướng gradient; Adam chia bước theo moment bậc hai. So lr tốt nhất trong phạm vi đã thử, không khẳng định tối ưu toàn cục.',
 'loss':'CE có gradient logit p-y. MSE ở đây dùng logit thô, mean trên 7 tọa độ: gradient 2(z-onehot)/7; không so trị số loss giữa CE và MSE.',
 'hparam':'Tăng lr làm bước cập nhật lớn hơn; cùng 20 epoch và batch 512 giữ số bước như nhau.',
 'dropout':'Dropout tăng nhiễu lúc train; cả train loss và val loss đều đo eval mode. Gap cuối dương là dấu hiệu cần xét, không đủ để kết luận mọi regularization có lợi.',
 'clipping':'Global L2 clipping nhân gradient với min(1,c/||g||); giữ hướng nhưng giảm bước. Nó không sửa được nhãn sai hoặc lr quá nhỏ.',
 'amp':'Autocast BF16 chỉ bọc forward/loss, tham số và optimizer vẫn FP32. CPU không có CUDA nên không đo FP16 hay CUDA peak memory; thời gian phụ thuộc phần cứng.',
 'init':'ReLU cần phá đối xứng: zeros làm gradient các lớp ẩn bằng 0. He dùng Var=2/fan_in; Xavier dùng 2/(fan_in+fan_out). Stats đo sau từng Linear.'}

def health_checks(data,out):
    set_seed(1)
    model=MLP().to(data['X_tr'].device)
    assert count_params(model)==EXPECTED_PARAMS[(256,128)]
    assert model(data['X_tr'][:8]).shape==(8,7)
    step0=evaluate(model,data['X_val'],data['y_val'])['loss']
    loss=compute_loss(model(data['X_tr'][:512]),data['y_tr'][:512],'ce')
    loss.backward()
    grads={n:float(p.grad.norm()) for n,p in model.named_parameters()}
    assert all(v>0 and np.isfinite(v) for v in grads.values())
    model=MLP().to(data['X_tr'].device)
    opt=torch.optim.Adam(model.parameters(),lr=.01)
    curve=[]
    for _ in range(400):
        model.train();opt.zero_grad(set_to_none=True)
        loss=compute_loss(model(data['X_tr'][:20]),data['y_tr'][:20],'ce')
        loss.backward();opt.step();curve.append(float(loss.detach()))
    assert curve[-1]<.02
    fig,ax=plt.subplots(figsize=(7,4));ax.plot(range(1,401),curve)
    ax.set(xlabel='Update',ylabel='CE loss',title='Overfit 20 training samples (no regularization)');ax.set_yscale('log')
    fig.savefig(out/'figures/health_overfit20.png',dpi=130,bbox_inches='tight');plt.close(fig)
    health=dict(params=count_params(model),logits_shape=[8,7],step0_loss=step0,
                ln7=float(np.log(7)),overfit20_final_loss=curve[-1],gradient_norms=grads,
                train_numeric_mean=data['X_tr'][:,:10].mean(0).cpu().tolist(),
                train_numeric_std=data['X_tr'][:,:10].std(0,unbiased=False).cpu().tolist(),
                majority_val_acc=float((data['y_val']==1).float().mean()),
                torch_version=torch.__version__,python_version=platform.python_version(),device=str(data['X_tr'].device))
    (out/'results/health_checks.json').write_text(json.dumps(health,indent=2))
    print(json.dumps(health,indent=2))
    return health

def configs(group,clip_norm=None):
    variants={
      'baseline':[(f'base-s{s}',dict(seed=s)) for s in (1,2,3)],
      'optimizer':[('opt-sgdm-lr01',dict(lr=.01)),('opt-sgdm-lr10',dict(lr=.1)),
                   ('opt-adam-lr001',dict(optimizer='adam',lr=.001)),('opt-adam-lr003',dict(optimizer='adam',lr=.003))],
      'loss':[('loss-mse',dict(loss='mse'))],
      'hparam':[('lr-high',dict(lr=.5))],
      'dropout':[('dropout-q02',dict(dropout=.2))],
      'clipping':[('clip-normal',dict(clip_norm=clip_norm)),('clip-high',dict(lr=.5,clip_norm=clip_norm))],
      'amp':[('amp-bf16',dict(precision='bf16'))],
      'init':[('init-xavier',dict(init='xavier')),('init-zeros',dict(init='zeros'))]}
    return [{**DEFAULT_CFG,**change,'exp_id':exp_id,'group':group,
             'description':PREDICTIONS[group]} for exp_id,change in variants[group]]

def run_group(group,data,out,all_results,clip_norm=None):
    for cfg in configs(group,clip_norm):
        result=run_experiment(cfg,data)
        save_result(result,out/'results')
        plot_run(result,out/f"figures/{cfg['exp_id']}.png")
        all_results.append(result)
    reference=next(r for r in all_results if r['cfg']['exp_id']=='base-s1')
    subset=[r for r in all_results if r['cfg']['group']==group]
    if group!='baseline': subset=[reference]+subset
    if group=='clipping': subset.append(next(r for r in all_results if r['cfg']['exp_id']=='lr-high'))
    plot_compare(subset,'val_macro_f1',out/f'figures/compare_{group}.png',group)
    return subset

def describe_group(group,results,baseline,noise):
    print('Cơ chế:',MECHANISMS[group])
    for r in results:
        s=r['summary']; delta=s['val_macro_f1']-baseline['summary']['val_macro_f1']
        print(f"{r['cfg']['exp_id']}: F1={s['val_macro_f1']:.6f}, delta vs base-s1={delta:+.6f}, |delta| {'>' if abs(delta)>noise else '<='} 2σ={noise:.6f}; best epoch={s['best_epoch']}, gap cuối={s['final_val_loss']-s['final_train_loss']:.6f}, {s['time_per_epoch_s']:.3f}s/epoch")

def official_scores(result,data,out,repo,baseline=False):
    pred=out/('predictions_baseline.csv' if baseline else 'predictions_eval.csv')
    dest=out/('eval_baseline.json' if baseline else 'eval_result.json')
    final_eval(result['cfg'],result,data,str(pred))
    subprocess.run([sys.executable,str(repo/'scripts/evaluate.py'),'--pred',str(pred),'--out',str(dest),
        '--data',str(repo/'data/covtype.csv.gz'),'--meta',str(repo/'data/split_metadata.csv')],cwd=repo,check=True)
    return json.loads(dest.read_text())

def finish(results,data,health,out,repo):
    baseline=next(r for r in results if r['cfg']['exp_id']=='base-s1')
    # Configuration selection is completed and persisted BEFORE any eval scoring.
    candidates=[r for r in results if r['cfg']['seed']==1 and not r['summary']['diverged']]
    final=max(candidates,key=lambda r:r['summary']['val_macro_f1'])
    choice=dict(exp_id=final['cfg']['exp_id'],cfg=final['cfg'],
                selection='maximum val macro-F1 among seed-1 runs; checkpoint minimum val loss',
                val_macro_f1=final['summary']['val_macro_f1'])
    (out/'results/selection.json').write_text(json.dumps(choice,indent=2))
    print('Selected BEFORE eval:',choice)
    bscore=official_scores(baseline,data,out,repo,True)
    fscore=bscore if final is baseline else official_scores(final,data,out,repo)
    if final is baseline:
        import shutil
        shutil.copyfile(out/'predictions_baseline.csv',out/'predictions_eval.csv')
        shutil.copyfile(out/'eval_baseline.json',out/'eval_result.json')
    scores={baseline['cfg']['exp_id']:bscore,final['cfg']['exp_id']:fscore}
    rows=[to_row(r,scores.get(r['cfg']['exp_id']),MECHANISMS[r['cfg']['group']]) for r in results]
    # Recreate workbook from bundled template, making submission portable without templates/.
    write_xlsx(rows,Path(__file__).parent/'experiment_table_template.xlsx',out/'experiments.xlsx')
    cm=np.array(fscore['confusion_matrix'])
    fig,ax=plt.subplots(figsize=(7,6));im=ax.imshow(cm,cmap='Blues')
    for i in range(7):
        for j in range(7):ax.text(j,i,str(cm[i,j]),ha='center',va='center',fontsize=8,color='white' if cm[i,j]>cm.max()/2 else 'black')
    ax.set(xlabel='Predicted class',ylabel='True class',xticks=range(7),yticks=range(7),title='Eval confusion matrix');fig.colorbar(im,ax=ax)
    fig.savefig(out/'figures/eval_confusion.png',dpi=140,bbox_inches='tight');plt.close(fig)
    write_report(results,health,baseline,final,bscore,fscore,out)
    augment_workbook(results,health,bscore,fscore,out)
    return choice,bscore,fscore

def write_report(results,health,baseline,final,bscore,fscore,out):
    seeds=[r['summary'] for r in results if r['cfg']['group']=='baseline']
    mean=float(np.mean([s['val_macro_f1'] for s in seeds]));std=float(np.std([s['val_macro_f1'] for s in seeds],ddof=1));noise=2*std
    acc_mean=float(np.mean([s['val_acc'] for s in seeds]));acc_std=float(np.std([s['val_acc'] for s in seeds],ddof=1))
    base=baseline['summary']['val_macro_f1']
    lines=[f'# Báo cáo Lab Day 1 — Nguyễn Văn Giáp — 2A202602903',
      '\n## 1. Thiết lập',
      f"Python {health['python_version']}, PyTorch {health['torch_version']}, CPU (2 threads). Forest CoverType: 464 809 train / 116 203 eval theo metadata cố định; validation phân tầng 20%, seed 42: 371 847 train / 92 962 val. Chỉ fit mean/std 10 cột số trên train; 44 cột one-hot giữ nguyên.",
      'MLP tự định nghĩa 54→256→128→7, ReLU, bias, 47 879 tham số. Baseline CE, SGD+momentum 0.9, lr=.05, batch=512, 20 epoch, He ở mọi Linear, bias=0. Mọi thí nghiệm cùng split, 20 epoch; batch cuối vẫn dùng. Adam đổi optimizer và lr để quét lr riêng; so optimizer ở lr tốt nhất đã thử. clip-high đổi lr + clipping so baseline, nhưng chỉ đổi clipping khi so với lr-high. Các biến thể khác đổi một yếu tố. Train loss đo eval mode trên 50 000 mẫu đầu cố định. Checkpoint chọn theo val loss thấp nhất; cấu hình cuối chọn val macro-F1 của checkpoint, chỉ xét seed 1 để không chọn seed may mắn. Chạy đủ 7 chủ đề, mixed precision dùng BF16 CPU.',
      '\n## 2. Kiểm tra ban đầu và độ nhiễu',
      '| Kiểm tra | Kết quả |\n|---|---|',
      f"| Số tham số / logits | {health['params']} / (8, 7) |",
      f"| Loss bước 0 (`base-s1`) / ln 7 | {health['step0_loss']:.6f} / {health['ln7']:.6f} |",
      f"| Quá khớp 20 mẫu, loss cuối sau 400 bước | {health['overfit20_final_loss']:.8f} |",
      '| Gradient tới mọi tham số | Có, tất cả norm > 0 ở He |',
      f'| Baseline 3 seed, val acc TB ± σ | {acc_mean:.6f} ± {acc_std:.6f} |',
      f'| Baseline 3 seed, val macro-F1 TB ± σ | {mean:.6f} ± {std:.6f} |',
      f'| Ngưỡng 2σ | {noise:.6f} |',
      f"Loss đầu khác ln 7 vì He ở lớp ra tạo logit không đồng đều (không buộc logit=0). Đây không tự động là lỗi: gradient có mặt và kiểm tra 20 mẫu hội tụ. Accuracy đa số trên val là {health['majority_val_acc']:.6f}; baseline vượt mốc này. Health checks nằm ở `results/health_checks.json`; seed và cấu hình nằm trong bảng.",
      '![](figures/health_overfit20.png)',
      '\n## 3. Kết quả theo chủ đề',
      f'Các điểm dưới đây là val tại epoch có val loss thấp nhất. Chênh lệch dùng so với `base-s1`; |Δ| > {noise:.6f} chỉ là vượt ngưỡng nhiễu thực nghiệm, không phải kiểm định thống kê. Ba seed chỉ đo nhiễu baseline, mỗi biến thể mới có một seed.']
    for group in PREDICTIONS:
        subset=[r for r in results if r['cfg']['group']==group]
        lines += [f'\n### {group}',f'**Dự đoán trước:** {PREDICTIONS[group]}',
          '| exp_id | lr | best epoch | val macro-F1 | Δ vs base-s1 | >2σ? | gap loss cuối | s/epoch |',
          '|---|---:|---:|---:|---:|---|---:|---:|']
        for r in subset:
            s=r['summary'];d=s['val_macro_f1']-base
            lines.append(f"| {r['cfg']['exp_id']} | {r['cfg']['lr']} | {s['best_epoch']} | {s['val_macro_f1']:.6f} | {d:+.6f} | {'Có' if abs(d)>noise else 'Chưa'} | {s['final_val_loss']-s['final_train_loss']:.6f} | {s['time_per_epoch_s']:.3f} |")
        verdict = '; '.join(f"`{r['cfg']['exp_id']}` {'bằng' if r['summary']['val_macro_f1'] == base else ('cao hơn' if r['summary']['val_macro_f1'] > base else 'thấp hơn')} baseline {abs(r['summary']['val_macro_f1']-base):.6f}" for r in subset)
        lines += [f'**Đối chiếu:** {verdict}. Nếu chênh lệch không vượt 2σ, chưa kết luận khác biệt ổn định. **Cơ chế:** {MECHANISMS[group]}',f'![](figures/compare_{group}.png)']
        if group=='baseline':
            h=baseline['history']
            lines.append(f"Baseline train loss {h['train_loss'][0]:.6f} → {h['train_loss'][-1]:.6f}; val loss {h['val_loss'][0]:.6f} → {h['val_loss'][-1]:.6f}. Gap cuối {h['val_loss'][-1]-h['train_loss'][-1]:.6f}; cả hai loss giảm nhưng có dao động, nên 20 epoch chưa cho thấy quá khớp kéo dài.")
        if group=='dropout':
            r=subset[0];h=r['history'];bh=baseline['history']
            lines.append(f"Gap cuối giảm từ {bh['val_loss'][-1]-bh['train_loss'][-1]:.6f} xuống {h['val_loss'][-1]-h['train_loss'][-1]:.6f}, nhưng F1 thay đổi {r['summary']['val_macro_f1']-base:+.6f}. Giảm gap không đồng nghĩa tăng khả năng tổng quát; thêm nhiễu có thể làm chưa khớp trong cùng ngân sách 20 epoch.")
        if group=='optimizer':
            bests=[]
            for optimizer in ('sgd_momentum','adam'):
                pool=[r for r in results if r['cfg']['optimizer']==optimizer and r['cfg']['group'] in ('baseline','optimizer') and r['cfg']['seed']==1]
                bests.append(max(pool,key=lambda r:r['summary']['val_macro_f1']))
            lines.append('Ở lr tốt nhất đã thử: '+ '; '.join(f"{r['cfg']['optimizer']} `{r['cfg']['exp_id']}` lr={r['cfg']['lr']}, F1={r['summary']['val_macro_f1']:.6f}" for r in bests)+'.')
        if group=='clipping':
            for r in subset:
                lines.append(f"`{r['cfg']['exp_id']}`: c={r['cfg']['clip_norm']:.6f}; tỷ lệ lô bị clip trung bình {np.mean(r['history']['clip_fraction']):.4f}; norm lớn nhất trước clip {max(r['history']['grad_norm_max']):.4f}.")
        if group=='clipping':
            high=next(r for r in results if r['cfg']['exp_id']=='lr-high')
            clipped=next(r for r in subset if r['cfg']['exp_id']=='clip-high')
            delta=clipped['summary']['val_macro_f1']-high['summary']['val_macro_f1']
            lines.append(f"So cặp lr cao: clip-high F1={clipped['summary']['val_macro_f1']:.6f} vs lr-high {high['summary']['val_macro_f1']:.6f}, Δ={delta:+.6f}; {'vượt' if abs(delta)>noise else 'chưa vượt'} 2σ. Lr-high không phân kỳ; không gọi clipping là cứu phân kỳ trong lần chạy này.")
        if group=='amp':
            r=subset[0]
            lines.append(f"BF16/FP32 time ratio = {r['summary']['time_per_epoch_s']/baseline['summary']['time_per_epoch_s']:.3f}; BF16 {'nhanh hơn' if r['summary']['time_per_epoch_s']<baseline['summary']['time_per_epoch_s'] else 'chậm hơn'} trên CPU đã đo. Các thời gian ghi ở cột time_per_epoch_s; không có phép đo bộ nhớ CUDA.")
        if group=='init':
            lines.append('Activation std sau từng Linear / loss bước 0: '+ '; '.join(f"`{r['cfg']['exp_id']}` {np.round(r['summary']['activation_std'],6).tolist()} / {r['summary']['step0_loss']:.6f}" for r in [baseline]+subset)+'.')
    sgdm=max([baseline]+[r for r in results if r['cfg']['exp_id'].startswith('opt-sgdm')],key=lambda r:r['summary']['val_macro_f1'])
    adam=max([r for r in results if r['cfg']['optimizer']=='adam'],key=lambda r:r['summary']['val_macro_f1'])
    lines += ['\n## 4. Đánh giá cuối trên eval',
      f"Chọn `{final['cfg']['exp_id']}` bằng val trước khi gọi script chấm; lưu quyết định trong `results/selection.json`. Không điều chỉnh cấu hình sau khi xem eval.",
      '| Cấu hình | seed | val macro-F1 | eval macro-F1 | eval accuracy |\n|---|---:|---:|---:|---:|']
    for r,score in ((baseline,bscore),(final,fscore)):
        lines.append(f"| {r['cfg']['exp_id']} | {r['cfg']['seed']} | {r['summary']['val_macro_f1']:.6f} | {score['macro_f1']:.6f} | {score['accuracy']:.6f} |")
    improvement=fscore['macro_f1']-bscore['macro_f1']
    lines += [f"Cải thiện eval = {improvement:+.6f}; độ lệch eval–val của cấu hình cuối = {fscore['macro_f1']-final['summary']['val_macro_f1']:+.6f}. {'Vượt' if abs(improvement)>noise else 'Không vượt'} ngưỡng 2σ val. Không đo σ eval vì chỉ chấm baseline seed 1 và cấu hình cuối; không coi σ val là ước lượng σ eval.",
      '\n### Phân tích lỗi theo lớp',
      '| Lớp | support | precision | recall | F1 |\n|---|---:|---:|---:|---:|']
    for c in fscore['per_class']:
        lines.append(f"| {c['cls']} | {c['support']} | {c['precision']:.6f} | {c['recall']:.6f} | {c['f1']:.6f} |")
    worst=min(fscore['per_class'],key=lambda c:c['f1']);cm=np.array(fscore['confusion_matrix']);i=worst['cls'];row=cm[i].copy();row[i]=0;j=int(row.argmax())
    lines += [f"Lớp khó nhất {i}: F1={worst['f1']:.6f}, support={worst['support']}; nhầm nhiều nhất sang lớp {j} ({row[j]} mẫu). Mất cân bằng có thể giảm số bước học cho lớp hiếm; đặc trưng địa hình giữa lớp có thể chồng lấn, nhưng đây là giả thuyết chưa kiểm tra trực tiếp. Lần tiếp theo thử weighted CE và xem recall lớp này trên val.",
      '![](figures/eval_confusion.png)',
      '\n## 5. Trả lời câu hỏi dẫn dắt',
      f"1. SGD+momentum tốt nhất đã thử: `{sgdm['cfg']['exp_id']}` lr={sgdm['cfg']['lr']}, val F1={sgdm['summary']['val_macro_f1']:.6f}; Adam: `{adam['cfg']['exp_id']}` lr={adam['cfg']['lr']}, F1={adam['summary']['val_macro_f1']:.6f}. Khi không chỉnh lr có thể đảo thứ hạng; không suy rộng từ cùng một lr cho mọi optimizer.",
      '2. Dropout chỉ có lợi khi giảm overfit đủ bù việc giảm năng lực học; xem gap cuối và F1 ở mục dropout, không mặc định q càng cao càng tốt.',
      '3. Clipping hạn chế độ lớn bước khi gradient lớn; tỷ lệ clip và max norm ở mục clipping cho biết nó thực sự kích hoạt. Có thể giảm gai mà vẫn không tăng F1.',
      '4. BF16 được đo trực tiếp với FP32 ở mục amp. Tốc độ phụ thuộc kernel CPU và overhead autocast của MLP nhỏ; kết quả này không trả lời được tốc độ trên GPU FP16.',
      '5. Zeros khiến lớp ẩn ReLU bằng 0, gradient trọng số bằng 0; bias lớp cuối học phân bố nhãn. He bù việc ReLU làm mất một phần phương sai; Xavier cân đối fan-in/fan-out.',
      '6. Nếu loss không giảm sau 2 000 bước: (a) kiểm tra shape, dtype, nhãn 0..6, chuẩn hoá chỉ từ train và loss bước 0 để phát hiện dữ liệu/logit bất thường; (b) quá khớp 20 mẫu không regularization để kiểm tra pipeline update; (c) xem gradient từng lớp và norm trước clip, lr, zero_grad/backward/step để phát hiện gradient chết, bước quá nhỏ hoặc nổ gradient.',
      '\n## 6. Hạn chế và điều bất ngờ',
      'Chỉ ba seed baseline, một seed mỗi biến thể; so sánh vượt 2σ chưa phải bằng chứng thống kê đầy đủ. Lr được quét trong phạm vi nhỏ (SGD+momentum .01/.05/.1, Adam .001/.003), không tìm tối ưu toàn cục. MSE dùng lr baseline nên kết luận về loss phụ thuộc lr. BF16 CPU không thay cho thí nghiệm FP16 GPU; không có số đo CUDA memory. Health loss đầu cao hơn ln 7 là do logit khởi tạo, không ép kết quả về mốc lý thuyết. Chỉ chạy 20 epoch nên mô hình có thể còn chưa khớp; quan sát F1/loss theo epoch để đánh giá. σ từ checkpoint chọn trên val có thể lạc quan. Train loss dùng tập con cố định, không toàn bộ train.',
      '\n## 7. Phụ lục',
      'Bài nộp: REPORT.md, experiments.xlsx (Legend/Experiments/Seeds/Summary), predictions_eval.csv, eval_result.json, eval_baseline.json, predictions_baseline.csv, figures/, results/, code/ gồm notebook và tất cả module. Template bảng đi kèm trong code/ để chạy lại khi chỉ có bài nộp + data/ + scripts/. Không nộp data hay trọng số.',
      f"{len(results)} thí nghiệm, mỗi thí nghiệm một JSON và ảnh có đúng exp_id; tổng thời gian epoch đo được {sum(sum(r['history']['epoch_time_s']) for r in results):.1f}s. Chạy lại bằng notebook Restart & Run All hoặc `python code/run_lab.py` từ thư mục nộp."]
    # Keep tables contiguous while separating headings, paragraphs and figures.
    rendered=[]
    for line in lines:
        line=line.strip()
        if rendered:
            rendered.append('\n' if line.startswith('|') and rendered[-1].startswith('|') else '\n\n')
        rendered.append(line)
    (out/'REPORT.md').write_text(''.join(rendered)+'\n',encoding='utf-8')


def augment_workbook(results,health,bscore,fscore,out):
    """Record diagnostics and class analysis in Summary for numerical traceability."""
    import openpyxl
    from openpyxl.styles import Font
    wb=openpyxl.load_workbook(out/'experiments.xlsx')
    ws=wb['Summary']
    records=[('Diagnostics / supplementary measurements','Value','Source'),
      ('params',health['params'],'health_checks.json'),
      ('shape logits',str(health['logits_shape']),'health_checks.json'),
      ('ln7',health['ln7'],'reference'),
      ('overfit20_final_loss',health['overfit20_final_loss'],'health_checks.json'),
      ('majority_val_acc',health['majority_val_acc'],'health_checks.json'),
      ('total_epoch_time_s',sum(sum(r['history']['epoch_time_s']) for r in results),'all histories')]
    for n,v in health['gradient_norms'].items():records.append(('gradient '+n,v,'health_checks.json'))
    for r in results:
        exp=r['cfg']['exp_id'];h=r['history'];s=r['summary']
        records += [(exp+' grad_norm_p90',s['grad_norm_p90'],exp),
                    (exp+' max_grad_norm',max(h['grad_norm_max']),exp),
                    (exp+' mean_clip_fraction',float(np.mean(h['clip_fraction'])),exp)]
        for i,v in enumerate(s['activation_std'],1):records.append((exp+f' activation_std_linear{i}',v,exp))
    row=17
    for record in records:
        for col,v in enumerate(record,1):ws.cell(row,col,v)
        row+=1
    for label,score in (('baseline',bscore),('final',fscore)):
        row+=1
        for col,v in enumerate((label+' eval class','support','precision','recall','F1'),1):ws.cell(row,col,v)
        row+=1
        for c in score['per_class']:
            for col,key in enumerate(('cls','support','precision','recall','f1'),1):ws.cell(row,col,c[key])
            row+=1
        ws.cell(row,1,label+' confusion matrix (true rows / predicted columns)');row+=1
        for counts in score['confusion_matrix']:
            for col,v in enumerate(counts,1):ws.cell(row,col,v)
            row+=1
    for cell in ws[17]:cell.font=Font(bold=True)
    # Numeric aggregates in Seeds and Summary support readers without recalculation;
    # original experiment formula columns are retained.
    wb.save(out/'experiments.xlsx')
    from results_table import cache_formula_values
    cache_formula_values(out/'experiments.xlsx',[to_row(r) for r in results])
