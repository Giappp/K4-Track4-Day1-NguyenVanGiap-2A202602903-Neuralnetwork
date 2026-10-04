"""JSON histories and workbook export preserving the four template sheets."""
import json
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.formula.translate import Translator

def save_result(result,results_dir='../results'):
    path=Path(results_dir)/f"{result['cfg']['exp_id']}.json"
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps({k:result[k] for k in ('cfg','history','summary')},indent=2,allow_nan=False))
    return str(path)

def load_results(results_dir='../results'):
    return [json.loads(p.read_text()) for p in sorted(Path(results_dir).glob('*.json')) if p.name not in ('health_checks.json','selection.json')]

def to_row(result,eval_scores=None,notes=''):
    row={**result['cfg'],**result['summary'], 'figure_file':f"figures/{result['cfg']['exp_id']}.png",'notes':notes}
    row['hidden']='-'.join(map(str,row['hidden']))
    row['clip_norm']='none' if row['clip_norm'] is None else row['clip_norm']
    if eval_scores:
        row.update(eval_acc=eval_scores['accuracy'],eval_macro_f1=eval_scores['macro_f1'])
    if row['peak_mem_MB'] is None:
        row['notes']+='; CPU: CUDA peak memory unavailable; train loss on fixed 50,000 samples'
    return row

def write_xlsx(rows,template_path,out_path):
    wb=openpyxl.load_workbook(template_path)
    ws=wb['Experiments']
    headers=[c.value for c in ws[1]]
    formulas={c.column:c.value for c in ws[2] if c.data_type=='f'}
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if cell.column not in formulas:
                cell.value=None
    for i,row in enumerate(rows,2):
        for j,key in enumerate(headers,1):
            if j not in formulas:
                ws.cell(i,j,row.get(key))
            else:
                origin=ws.cell(2,j).coordinate
                ws.cell(i,j,Translator(formulas[j],origin=origin).translate_formula(ws.cell(i,j).coordinate))
    seeds=[r for r in rows if r['group']=='baseline']
    ss=wb['Seeds']
    for i in range(2,7):
        ss.cell(i,1,seeds[i-2]['exp_id'] if i-2<len(seeds) else '')
    for i,r in enumerate(seeds,2):
        for j,key in enumerate(('val_acc','val_macro_f1','best_val_loss'),2):
            ss.cell(i,j,r[key])
    # Numeric aggregates also readable without Excel formula recalculation.
    for j,key in enumerate(('val_acc','val_macro_f1','best_val_loss'),2):
        vals=[r[key] for r in seeds]
        ss.cell(8,j,float(np.mean(vals)))
        ss.cell(9,j,float(np.std(vals,ddof=1)) if len(vals)>1 else 0.)
        ss.cell(10,j,2*float(np.std(vals,ddof=1)) if len(vals)>1 else 0.)
    su=wb['Summary']
    for i in range(2,12):
        group=su.cell(i,1).value
        group_rows=[r for r in rows if r['group']==group]
        for j,v in enumerate((len(group_rows),len(group_rows),max((r['val_macro_f1'] for r in group_rows),default=None),
                               min((r['val_macro_f1'] for r in group_rows),default=None),
                               max((r['val_acc'] for r in group_rows),default=None),('—' if group in ('baseline','final','other') else ('Có' if group_rows else 'Chưa'))),2):
            su.cell(i,j,v)
        su.cell(i,8,'; '.join(r['exp_id'] for r in group_rows) if group_rows else 'Không chạy trên phần cứng hiện tại')
    Path(out_path).parent.mkdir(parents=True,exist_ok=True)
    wb.save(out_path)


def cache_formula_values(path, rows):
    """Cache the template's derived cells while retaining their Excel formulas.

    openpyxl does not calculate formulas. We evaluate this template's four known
    derived columns from source records, so Excel and data_only readers agree.
    """
    import tempfile
    import zipfile
    import xml.etree.ElementTree as ET
    ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('',ns)
    wb=openpyxl.load_workbook(path,data_only=False)
    by_id={r['exp_id']:r for r in rows}
    seeds=[r for r in rows if r['group']=='baseline']
    mean=float(np.mean([r['val_macro_f1'] for r in seeds]))
    noise=2*float(np.std([r['val_macro_f1'] for r in seeds],ddof=1))
    caches={}
    for i,name in enumerate(wb.sheetnames,1):
        sheet=wb[name];values={}
        for cells in sheet:
            for c in cells:
                if c.data_type!='f':continue
                value=''
                if name=='Experiments':
                    record=by_id.get(sheet.cell(c.row,1).value)
                    if record:
                        if c.column==30:value=record['step0_loss']-float(np.log(7))
                        elif c.column==31:value=record['final_val_loss']-record['final_train_loss']
                        elif c.column==32:value=record['val_macro_f1']-mean
                        elif c.column==33:value='Có' if abs(record['val_macro_f1']-mean)>noise else 'Không'
                elif name=='Seeds' and 2<=c.row<=6:
                    record=by_id.get(sheet.cell(c.row,1).value)
                    if record:value=record[{2:'val_acc',3:'val_macro_f1',4:'best_val_loss'}[c.column]]
                elif name=='Summary' and c.coordinate=='D13':
                    value=sum(sheet.cell(r,7).value=='Có' for r in range(3,10))
                values[c.coordinate]=value
        caches[f'xl/worksheets/sheet{i}.xml']=values
    with zipfile.ZipFile(path) as zin:
        contents={name:zin.read(name) for name in zin.namelist()}
    for name,values in caches.items():
        if not values:continue
        root=ET.fromstring(contents[name])
        for cell in root.iter(f'{{{ns}}}c'):
            ref=cell.attrib.get('r')
            if ref not in values:continue
            value=values[ref]
            vnode=cell.find(f'{{{ns}}}v')
            if vnode is None:vnode=ET.SubElement(cell,f'{{{ns}}}v')
            cell.attrib['t']='str' if isinstance(value,str) else 'n'
            vnode.text=str(value)
        contents[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with tempfile.NamedTemporaryFile(dir=Path(path).parent,suffix='.xlsx',delete=False) as f:
        temp=Path(f.name)
    try:
        with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as zout:
            for name,content in contents.items():zout.writestr(name,content)
        temp.replace(path)
    finally:
        temp.unlink(missing_ok=True)
